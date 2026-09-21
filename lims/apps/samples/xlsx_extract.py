# -*- coding: utf-8 -*-
"""巴西 PLANILHA DE ENVIO 汇总表（xlsx）— 无创 NIPT 行提取。

输入：巴西实验室发送的「PLANILHA DE ENVIO」汇总表，表内**混装**两类业务：
    - Test_Item == "NIPPT"          → 无创产前亲子鉴定（NIPPT 子系统负责）
    - Test_Item 以 "NIPT" 开头       → 无创产前筛查（本模块负责）

本模块只提取后者，登记为 NIPT 无创样本；NIPPT 行进 skipped 交回调用方显示，
其它未识别 Test_Item 进 errors（不猜测、不登记）。

输出契约与 samples/docx_extract.extract_brazil_docx 保持同构，
以便 register_from_pdf 的建样本 / 查重 / 回执 Excel 逻辑 100% 复用。

表结构：第 1-2 行双语言表头（英文 / 葡语），数据自第 3 行起；按列位置取值。
"""
import datetime
import io
import re

import openpyxl

# ── Test_Item 档位 → (panel code, Sample.test_option) ──
# 新增档位只需在此加一行；未列出的 NIPT 档位会进 errors 让人工确认，不猜测。
PANEL_MAP = {
    "BÁSICO": ("NIPT", "Basic"),
    "BASICO": ("NIPT", "Basic"),
    "EXPANDIDO": ("NIPT_PLUS", "Plus"),
}

# 葡语备注 → 中文（未识别保留原文）
NOTE_TRANSLATIONS = [
    ("SEGUNDO SUPOSTO PAI", "第二位疑父"),
    ("PRIMEIRO SUPOSTO PAI", "第一位疑父"),
    ("RECOLETA GESTANTE SOLICITADA", "已要求孕妇重采"),
    ("AGUARDANDO AMOSTRA GESTANTE", "等待孕妇样本"),
    ("AMOSTRA DO SUPOSTO PAI", "疑父样本已到"),
    ("AMOSTRA DA GESTANTE", "孕妇样本已到"),
]

DATE_RE = re.compile(r"(\d{1,2})[/\-](\d{1,2})[/\-](\d{4})")

# 列索引（0-based）
COL_TEST_ITEM = 0
COL_SEQ = 1
COL_CLIENT_CODE = 2
COL_CLIENT_NAME = 3
COL_FEMALE_NAME = 4
COL_RG_F = 5
COL_DOB = 11
COL_GEST_WEEK = 12
COL_PRICE = 13
COL_SINAL = 14
COL_BALANCE = 15
COL_COLL_DATE = 19
COL_COLL_PLACE = 20
COL_FEDEX = 24
COL_DUE_DATE = 27
COL_GENDER = 31
COL_REMARKS = 32


class NipptPlanilhaParseError(Exception):
    pass


def _to_text(v):
    """单元格 → 字符串（datetime / 数字兼容）"""
    if v is None:
        return ""
    if isinstance(v, datetime.datetime):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, datetime.date):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).replace("\xa0", " ").strip()


def _parse_date_dmy(v):
    """从字符串提取第一个 DD/MM/YYYY 或 DD-MM-YYYY；失败返回 None"""
    s = _to_text(v)
    if not s:
        return None
    m = DATE_RE.search(s.replace("\n", " "))
    if not m:
        return None
    d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
    try:
        return datetime.date(y, mo, d)
    except ValueError:
        return None


def _first_part(v, sep="|"):
    """双值 'A | B' 取第一个（母亲值）"""
    s = _to_text(v)
    if not s:
        return ""
    return s.split(sep)[0].strip()


def _clean_money(v):
    """'R$ 2.550,00' → '2.550,00'；'-'/空 → ''"""
    s = _to_text(v)
    if not s or s == "-":
        return ""
    return s.replace("R$", "").strip()


def _extract_gestational_week(v):
    """'9 SEMANAS' → 9；'8 SEMANAS 5 DIAS' → 8；失败 None"""
    s = _to_text(v)
    m = re.search(r"(\d{1,2})", s)
    if not m:
        return None
    w = int(m.group(1))
    return w if 1 <= w <= 45 else None


def _map_gender(g):
    """Gender(SEXAGEM) 列 Sim/Não → Yes/No（与 Word 通道输出格式一致）"""
    s = _to_text(g).lower().replace("\u00e3", "a").replace("-", "").strip()
    if s in ("sim", "yes", "s", "y"):
        return "Yes"
    if s in ("nao", "no", "n"):
        return "No"
    return ""


def _translate_notes(remarks):
    """葡语备注 → 中文（未识别保留原文），多条用 '；' 连接"""
    s = _to_text(remarks)
    if not s or s == "-":
        return ""
    parts = [p.strip() for p in re.split(r"[/;|]", s) if p.strip()]
    out = []
    for p in parts:
        up = p.upper()
        matched = next((cn for pt, cn in NOTE_TRANSLATIONS if pt in up), None)
        out.append(matched or p)
    return "；".join(out)


def _assign_sales_agent(seq, src_code):
    """与 docx 通道同规则：seq 纯数字 → sales，否则 → agent"""
    sales, agent = "", ""
    if seq and str(seq).strip():
        if str(seq).isdigit():
            sales = src_code
        else:
            agent = src_code
    return sales, agent


def _resolve_panel(test_item):
    """Test_Item → (panel_code, test_option)；未识别档位返回 (None, None)"""
    up = _to_text(test_item).upper()
    for key, mapped in PANEL_MAP.items():
        if key in up:
            return mapped
    return None, None


def extract_brazil_planilha(file_path, source="巴西万基"):
    """解析 PLANILHA DE ENVIO xlsx。

    返回 (infos, skipped, errors)：
        infos   : list[dict]，契约同 docx_extract.extract_brazil_docx
        skipped : list[dict]  — 识别出的 NIPPT 行（亲子子系统负责，非错误）
        errors  : list[dict]  — 未识别 Test_Item / 未识别档位 / 缺关键字段的行
    """
    try:
        with open(file_path, "rb") as fh:
            data = fh.read()
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as e:
        raise NipptPlanilhaParseError(f"无法读取 xlsx: {str(e)[:150]}")

    ws = wb[wb.sheetnames[0]] if wb.sheetnames else None
    if ws is None:
        wb.close()
        raise NipptPlanilhaParseError("xlsx 中没有工作表")

    infos, skipped, errors = [], [], []

    row_iter = ws.iter_rows(values_only=True)
    # 跳过前两行双语表头
    for _ in range(2):
        try:
            next(row_iter)
        except StopIteration:
            break

    for row_no, r in enumerate(row_iter, start=3):
        if not r:
            continue
        test_item = _to_text(r[COL_TEST_ITEM] if len(r) > COL_TEST_ITEM else "")
        client_code = _to_text(r[COL_CLIENT_CODE] if len(r) > COL_CLIENT_CODE else "")
        seq = _to_text(r[COL_SEQ] if len(r) > COL_SEQ else "")
        mother_name = _to_text(r[COL_FEMALE_NAME] if len(r) > COL_FEMALE_NAME else "")
        client_name = _to_text(r[COL_CLIENT_NAME] if len(r) > COL_CLIENT_NAME else "")

        # 空行 —— 全空则跳过
        if not test_item and not client_code and not mother_name and not client_name:
            continue

        # ── 分流 ──
        up = test_item.upper()
        if not up:
            errors.append({"row": row_no, "test_item": "", "reason": "Test_Item 为空"})
            continue
        if up == "NIPPT":
            # 亲子样本，由 NIPPT 子系统导入通道负责
            skipped.append({
                "row": row_no, "test_item": test_item,
                "seq": seq, "client_code": client_code,
                "reason": "NIPPT（亲子样本，请在 NIPPT 登记页导入）",
            })
            continue
        if not up.startswith("NIPT"):
            errors.append({
                "row": row_no, "test_item": test_item,
                "reason": "Test_Item 不在白名单（非 NIPT 项目，未登记）",
            })
            continue

        panel_code, test_option = _resolve_panel(test_item)
        if not panel_code:
            errors.append({
                "row": row_no, "test_item": test_item,
                "reason": "NIPT 档位未识别，未登记（请在 PANEL_MAP 中补充映射）",
            })
            continue

        if not client_code:
            errors.append({
                "row": row_no, "test_item": test_item,
                "reason": "缺少 Client Code（外部编号/查重键）",
            })
            continue

        # ── 字段提取 ──
        dob_raw = _to_text(r[COL_DOB] if len(r) > COL_DOB else "")
        coll_date_raw = _to_text(r[COL_COLL_DATE] if len(r) > COL_COLL_DATE else "")
        gest_raw = _to_text(r[COL_GEST_WEEK] if len(r) > COL_GEST_WEEK else "")
        src_code = ""
        m = re.match(r"^VGBR([A-Z]{2,4})\d+", client_code)
        if m:
            src_code = m.group(1)
        sales, agent = _assign_sales_agent(seq, src_code)

        infos.append({
            "sample_source": source,
            "test_option": test_option,
            "panel_code": panel_code,
            "test_item": test_item,
            "external_id": client_code,
            "seq": seq,
            "sales": sales,
            "agent": agent,
            "patient_name": client_name or mother_name,
            "patient_dob_raw": dob_raw,
            "patient_dob": _parse_date_dmy(dob_raw),
            "id_card": _to_text(r[COL_RG_F] if len(r) > COL_RG_F else ""),
            "gestational_weeks_raw": gest_raw,
            "gestational_weeks": _extract_gestational_week(gest_raw),
            "age": None,
            "multiple_gestation": False,
            "ivf_status": False,
            "clinical_diagnosis": "",
            "ordering_facility": _first_part(r[COL_COLL_PLACE] if len(r) > COL_COLL_PLACE else ""),
            "ordering_physician": "",
            "collection_date_raw": coll_date_raw,
            "collection_date": _parse_date_dmy(_first_part(coll_date_raw)),
            "price": _clean_money(r[COL_PRICE] if len(r) > COL_PRICE else ""),
            "sinal": _clean_money(r[COL_SINAL] if len(r) > COL_SINAL else ""),
            "balance": _clean_money(r[COL_BALANCE] if len(r) > COL_BALANCE else ""),
            "fetal_gender": _map_gender(r[COL_GENDER] if len(r) > COL_GENDER else ""),
            "fedex_no": _to_text(r[COL_FEDEX] if len(r) > COL_FEDEX else ""),
            "report_due_date": _first_part(_to_text(r[COL_DUE_DATE] if len(r) > COL_DUE_DATE else "")),
            "remarks_raw": _to_text(r[COL_REMARKS] if len(r) > COL_REMARKS else ""),
            "remarks": _translate_notes(r[COL_REMARKS] if len(r) > COL_REMARKS else ""),
            "wd": "",
            "source_file": file_path.split("/")[-1],
        })

    wb.close()
    return infos, skipped, errors


if __name__ == "__main__":
    import json
    import sys

    infos, skipped, errors = extract_brazil_planilha(sys.argv[1])
    print(json.dumps({
        "infos": len(infos), "skipped": len(skipped), "errors": len(errors)
    }, ensure_ascii=False))
    for i in infos:
        print(" ", i["external_id"], "|", i["test_item"], "->", i["panel_code"],
              "|", i["patient_name"][:28], "|", i["gestational_weeks"],
              "|", i["collection_date"], "|", i["fetal_gender"], "|", i["report_due_date"])
    for s in skipped:
        print("  SKIP:", s["row"], s["test_item"], s["client_code"])
    for e in errors:
        print("  ERR :", e["row"], e["test_item"], e["reason"])
